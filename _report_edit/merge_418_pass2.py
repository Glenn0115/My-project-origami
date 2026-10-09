# -*- coding: utf-8 -*-
"""110 pass2：补充真实事实与论文式表述（全部基于仓库可查证据）。"""
import docx, sys, io, zipfile
from copy import deepcopy
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

PATH = r"Assets\从纸艺仿生到机械仿真设计应用研究报告_110_合并.docx"
d = docx.Document(PATH)
body = d.element.body

def para_text(p):
    return "".join(n.text or "" for n in p.iter(qn("w:t")))

def find_para(anchor):
    for p in body.findall(qn("w:p")):
        if para_text(p).startswith(anchor):
            return p
    raise SystemExit("NOT FOUND: " + anchor)

def set_text(p, text):
    rpr = None
    for r in p.findall(qn("w:r")):
        if r.find(qn("w:t")) is not None and r.find(qn("w:drawing")) is None:
            if r.find(qn("w:rPr")) is not None:
                rpr = deepcopy(r.find(qn("w:rPr")))
            break
    for child in list(p):
        if child.tag.split("}")[-1] in ("r", "hyperlink"):
            if child.find(qn("w:drawing")) is None and child.find(qn("w:pict")) is None:
                p.remove(child)
    run = OxmlElement("w:r")
    if rpr is not None:
        run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    ppr = p.find(qn("w:pPr"))
    p.insert(1 if ppr is not None else 0, run)

def clone_para(p):
    c = deepcopy(p)
    for r in list(c.findall(qn("w:r"))):
        c.remove(r)
    return c

body_template = find_para("仿生设计并不是对自然外形的简单复制")
appx3_heading = find_para("附录三 术语与完成状态说明")

def insert_after(anchor, texts, template=None):
    anchor_p = find_para(anchor)
    tmpl = template or body_template
    for t in texts:
        np_ = clone_para(tmpl)
        set_text(np_, t)
        anchor_p.addnext(np_)
        anchor_p = np_
    return anchor_p

n_ins = 0
# 2.1 仿生折纸分类（Fakhari [6]，真实文献）
insert_after("折纸工程综述指出", [
 "在仿生方向，近期综述研究开始更明确地把生物形态、行为与折纸工程联系起来。Fakhari 等建立了从动物、植物到 DNA 折纸的生物启发折纸分类框架，指出仿生折纸不仅服务于形态模仿，更服务于可持续材料、软体机器人与生物启发机构设计[6]。这说明当代仿生折纸研究的重点，已经从“像不像自然”转向“是否提取了自然对象的结构逻辑、运动机制与环境适应性”。"])
n_ins += 1
# 2.2 Kresling 动力学（双稳态/刚度调谐，仅作文献背景）
insert_after("对于设计研究，Kresling 的价值", [
 "已有研究还指出，Kresling 构型在一定几何参数区间内可呈现双稳态或多稳态响应，并可通过调整边数、相位角与高度比等参数实现刚度调谐[3-4]。这些动力学结果建立在工程材料与精密测量的基础上；本研究只用其几何构型与运动特性说明单元选择，不借用文献中的数值结果充当本样机的实验数据。"])
n_ins += 1
# 5.2 JSON 真实字段细则（依据仓库中 927_valley_drive.json 实际内容）
insert_after("工程中的模型数据以 JSON 保存", [
 "以单模块模型 927_valley_drive.json 为例：文件包含 name、description、vertices、faces、creases、connections、material 以及折痕宽度与颜色等字段。vertices 记录三维顶点坐标；faces 以 id、vertices、rigid、color 记录面片编号、顶点索引、刚体开关与颜色；creases 以 id、v1、v2、type、driveMode、actuatorGroup、restAngle、minAngle、maxAngle 记录折痕两端点、折痕类型、驱动方式与驱动器分组、初始角及允许角度范围。该模型共有 16 个顶点、16 个面片、32 条折痕，文件描述注明“仅驱动谷折、山折被动跟随”的手动滑条驱动方式；后续调试中又将驱动设置调整为仅驱动山折折痕（见 5.6 节）。模型文件在工程版本提交记录（5cd67c5“弹性模拟纸片”）中可查，同批提交还包括 927_elastic_flatten.json 弹性压平测试模型；当前工作区另保留 927_spatial.json、930_spatial.json 等空间导入测试模型。"])
n_ins += 1
# 6.2 应用情境细化（诚实口径）
insert_after("这些情境是设计推演", [
 "以教学演示情境为例，单模块折痕图可以直接用于课堂讲解山谷折与面片关系，Unity 画面适合演示折叠进度与状态切换，实体样机视频适合说明模块化连接后的整体效果；三者分别承担概念讲解、动态演示和实物参照的角色。以展示装置情境为例，Kresling 单元的轴向—扭转耦合适合表达可收缩的空间构件概念，但进入实际制作前仍需补充耐久性、连接强度与安全评估。上述表述均为情境层面的设计推演，不作为已完成应用的证据。"])
n_ins += 1
# 第十章 成果要点（诚实总结四点）
insert_after("报告同时提出正方形纸片", [
 "从证据组织角度看，本报告的价值主要体现在四点：一是仿生转译过程可回溯，自然运动特征与折线、模块参数之间的映射以表格与图注固定下来（表 2-1、表 4-2）；二是数字表达有据可查，展示对象限定为一个 Kresling 单模块，模型文件、场景与版本在文中注明（图 5-1、图 5-3、图 7-5）；三是实体证据与表述口径一致，三模块样机以照片与同源关键帧记录可见运动，未测量项如实标注（图 7-1 至图 7-4、表 7-1）；四是教学部分以可执行的任务与评价方案呈现，实验未实施即明确不纳入结果（第八章）。"])
n_ins += 1
# 附录三 扩充术语（全部与报告/工程实际用词一致）
last_bullet = insert_after("• 应用概念：根据结构特征提出的可能情境", [
 "• 折痕（Crease）：相邻面片的共享边，按类型区分为山折、谷折与边界折痕。",
 "• 山折／谷折：折痕折叠方向约定；山折为折线凸起、谷折为折线凹陷，图例见图 4-1。",
 "• 面片（Face）：折痕网络闭合区域对应的纸面单元，数字模型中以顶点索引与颜色记录。",
 "• 铰链（HingeJoint）：沿共享折痕建立的旋转约束，用于表达相邻面片绕折痕的转动。",
 "• 折叠进度（FoldProgress）：统一驱动全部有效折痕的 0—1 控制量，由滑条输入并映射为各铰链目标角。",
 "• 驱动器分组（actuatorGroup）：折痕主动驱动的分组字段；本报告单模块最终采用仅驱动山折的设置（见 5.6 节）。",
 "• 根面片：模拟中固定不动、作为其余面片折叠参照的基准面片。",
 "• 穿模／抖动：刚体与铰链模拟在参数不合适或碰撞复杂时出现的局部不稳定现象，属数字仿真边界，不解释为实体样机缺陷。",
 "• 共面／非共面 DXF：二维共面折线导入与保留三维坐标的空间导入两条路径，分别进入设计模式与模拟页面。",
 "• 双稳态（Bistability）：Kresling 构型在特定参数下可保持多个稳定状态的性质，本报告仅作为文献背景引用，未对样机开展稳态测量。",
 "• 中间层方法：位于仿生观察与工程分析之间，用于纸艺样机拆解、参数记录与数字转译的工作方法。",
 "• 证据链：从文献、折痕图、模型文件、运行画面到实体影像的可回溯记录体系，用于支撑每一步设计判断。",
])
n_ins += 12
# 新增附录四 工程文件与版本记录（真实 git 事实）
appx4 = clone_para(appx3_heading)
set_text(appx4, "附录四 工程文件与版本记录")
last_bullet.addnext(appx4)
anchor_p = appx4
for t in [
 "为便于核对，本报告涉及的关键工程文件与版本事实汇总如下。工程基于 Unity 2022.3.62f2c1，展示场景为 SampleScene。",
 "模型文件。单模块展示所用模型为 927_valley_drive.json，其字段结构与几何规模见 5.2 节。该文件与 927_elastic_flatten.json 弹性压平测试模型同批记录于工程版本提交 5cd67c5（“弹性模拟纸片”），当前工作区另保留 927_spatial.json、930_spatial.json 空间导入测试模型及 oldmodel 目录下的多组调试模型。",
 "调试记录。驱动与参数调试的关键提交包括：5b52f11（“硬件结构上应该没有什么问题了，就看 joint 阈值的设定了”）、0398c6f（“改 gap 可以让他动！！！！！”）、859bd79（“排查错误以及只驱动山折痕”）、77b1622（“gap 0.05 之后即可正常模拟”）。经上述调试，单模块采用仅驱动山折折痕、关节间隙 gap 取 0.05 的设置，可连续稳定完成折叠模拟（详见 5.6 节）。",
 "使用口径。本报告的数字成果只对应单模块展示画面；其他测试模型、空间导入模型与调试场景记录不构成本作品的成果，不与已完成展示混写。",
 "版本说明。报告写作期间工程持续更新，若后续补充新画面，应同步更新模型文件、场景与截图对应的版本记录。",
]:
    np_ = clone_para(body_template)
    set_text(np_, t)
    anchor_p.addnext(np_)
    anchor_p = np_
    n_ins += 1
# TOC：新增附录四条目
toc3 = find_para("附录三 术语与完成状态说明29")
toc4 = clone_para(toc3)
set_text(toc4, "附录四 工程文件与版本记录29")
toc3.addnext(toc4)
n_ins += 1

d.save(PATH)
print("inserted:", n_ins)

with zipfile.ZipFile(PATH) as z:
    bad = z.testzip()
print("zip ok" if bad is None else f"zip BAD: {bad}")

d2 = docx.Document(PATH)
paras = [p for p in d2.element.body.findall(qn("w:p"))]
total = sum(len("".join(n.text or "" for n in p.iter(qn("w:t")))) for p in paras)
tbl = sum(len(n.text or "") for tb in d2.tables for row in tb.rows for cell in row.cells for n in cell.paragraphs[0]._p.iter(qn("w:t")))
print(f"chars: body {total} + tables {tbl} = {total+tbl}")
