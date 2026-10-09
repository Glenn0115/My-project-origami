# -*- coding: utf-8 -*-
"""Patch: fix ch8 body heading (stray '23') and delete remaining stale placeholder lines."""
import docx, sys, io
from docx.oxml.ns import qn

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

PATH = r"Assets\艺仿生到机械仿真设计应用研究报告_109_修订.docx"
d = docx.Document(PATH)
body = d.element.body

def para_text(p):
    return "".join(n.text or "" for n in p.iter(qn("w:t")))

def set_text(p, text):
    from copy import deepcopy
    from docx.oxml import OxmlElement
    rpr = None
    for r in p.findall(qn("w:r")):
        if r.find(qn("w:t")) is not None and r.find(qn("w:drawing")) is None:
            rpr = r.find(qn("w:rPr"))
            if rpr is not None:
                rpr = deepcopy(rpr)
            break
    for child in list(p):
        if child.tag.split("}")[-1] in ("r", "hyperlink"):
            if child.find(qn("w:drawing")) is None and child.find(qn("w:pict")) is None:
                p.remove(child)
    run = OxmlElement("w:r")
    if rpr is not None:
        run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    ppr = p.find(qn("w:pPr"))
    p.insert(1 if ppr is not None else 0, run)

fixed = 0
for p in body.findall(qn("w:p")):
    txt = para_text(p)
    if txt == "第八章 软件教学方案与方案形成效率评价设计23":
        set_text(p, "第八章 软件教学方案与方案形成效率评价设计")
        fixed += 1
    if txt == "【此处插入对应图像，随后补齐图注信息】":
        p.getparent().remove(p)
        fixed += 1

d.save(PATH)
print("patched:", fixed)

# --- final verification dump ---
d2 = docx.Document(PATH)
paras = [p.text for p in d2.paragraphs]
out = []
keys = ["摘要", "8.5 方案形成效率实验设计", "8.6 结果处理方式说明", "8.7 本章小结",
        "5.5 数字仿真参数观察与定性验证", "5.6 技术边界", "5.7 本章小结",
        "7.5 综合设计评估", "9.2 对申报目标的审慎回应", "第十章 结论", "附录二"]
for i, t in enumerate(paras):
    if any(t.startswith(k) or t == k for k in keys):
        out.append(f"== {t}")
        for j in range(i + 1, min(i + 6, len(paras))):
            s = paras[j].strip()
            if s and (s.startswith(("表", "图")) or len(s) > 15):
                out.append("   " + s[:130])
            if s and not s.startswith(("表", "图")) and len(s) > 15:
                break
with open("_report_edit/_verify.txt", "w", encoding="utf-8") as f:
    f.write("\n".join(out))
print("verify dump written")
