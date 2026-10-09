# -*- coding: utf-8 -*-
"""110 pass3：半边结构面片生成、参考网格、方法特点（全部经仓库代码核实）。"""
import docx, sys, io, zipfile
from copy import deepcopy
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

PATH = r"Assets\从纸艺仿生到机械仿真设计应用研究报告_110_合并.docx"
d = docx.Document(PATH)
body = d.element.body

def para_text(p):
    return "".join(n.text or "" for n in p.iter(qn("w:t")))

def find_para(anchor):
    for p in body.findall(qn("w:p")):
        if para_text(p).startswith(anchor):
            return p
    raise SystemExit("NOT FOUND: " + anchor)

def set_text(p, text):
    rpr = None
    for r in p.findall(qn("w:r")):
        if r.find(qn("w:t")) is not None and r.find(qn("w:drawing")) is None:
            if r.find(qn("w:rPr")) is not None:
                rpr = deepcopy(r.find(qn("w:rPr")))
            break
    for child in list(p):
        if child.tag.split("}")[-1] in ("r", "hyperlink"):
            if child.find(qn("w:drawing")) is None and child.find(qn("w:pict")) is None:
                p.remove(child)
    run = OxmlElement("w:r")
    if rpr is not None:
        run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    ppr = p.find(qn("w:pPr"))
    p.insert(1 if ppr is not None else 0, run)

def clone_para(p):
    c = deepcopy(p)
    for r in list(c.findall(qn("w:r"))):
        c.remove(r)
    return c

body_template = find_para("仿生设计并不是对自然外形的简单复制")

def insert_after(anchor, texts):
    anchor_p = find_para(anchor)
    for t in texts:
        np_ = clone_para(body_template)
        set_text(np_, t)
        anchor_p.addnext(np_)
        anchor_p = np_

# 5.3 半边结构面片自动生成（HalfEdgeMesh.cs 实际算法）
insert_after("普通模型可由整体折叠进度控制", [
 "在面片生成环节，系统以半边结构组织折痕网络的拓扑关系。折痕编辑阶段主要维护“顶点—折痕”关系；保存时，OrigamiFaceGenerator 依据 HalfEdgeMesh 为每条折痕生成一对方向相反、互为 twin 的半边，并在每个顶点处按极角排序出射半边，把“下一出射半边的 twin”设为 next 指针。随后从任一半边出发沿 next 遍历，回到起点的闭合环即形成一个面片的顶点列表；未闭合或顶点数不足的环不参与面片生成。该机制使折痕图无需逐面手动标注即可自动推导面片结构，是二维折痕数据向三维面片转换的基础。",
])
# 5.4 参考网格（P191 网格吸附 + faces[].color 字段）
insert_after("设计模式与仿真模式围绕同一数据模型衔接", [
 "设计模式的参考网格为顶点定位提供规则间距的吸附基准（加点时鼠标落点自动吸附到网格），网格显示随编辑视角保持可读，便于在设计模式下完成对齐、测读与折线校对。面片颜色以模型文件中的 color 字段记录，用于区分不同纸面。",
])
# 9.3 方法特点四点（418 7.6 改写，诚实口径）
insert_after("从学科视角看", [
 "综合来看，本方法的特点可归纳为四点：可视化程度较高，结构关系、关键状态与参数变化能够在同一视域内被直观观察；试错成本较低，折痕图编辑与折叠模拟均适合快速试作与多轮比较；过程证据较完整，折痕图、JSON 模型文件、运行截图与实体影像可以形成连续记录链；设计与教学适配性较好，更适合设计前期、结构认知与课堂演示场景。上述特点是对方法适用性的概括，不表示已经完成工程性能验证。",
])

d.save(PATH)
with zipfile.ZipFile(PATH) as z:
    bad = z.testzip()
print("zip ok" if bad is None else f"zip BAD: {bad}")

d2 = docx.Document(PATH)
paras = [p for p in d2.element.body.findall(qn("w:p"))]
total = sum(len("".join(n.text or "" for n in p.iter(qn("w:t")))) for p in paras)
tbl = sum(len(n.text or "") for tb in d2.tables for row in tb.rows for cell in row.cells for n in cell.paragraphs[0]._p.iter(qn("w:t")))
print(f"chars: body {total} + tables {tbl} = {total+tbl}")
