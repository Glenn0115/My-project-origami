# -*- coding: utf-8 -*-
"""精确校验 110：图注紧跟在图片段之后；抽查关键区域；统计字数。"""
import docx, sys, io
from docx.oxml.ns import qn

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

PATH = r"Assets\从纸艺仿生到机械仿真设计应用研究报告_110_合并.docx"
d = docx.Document(PATH)
body = d.element.body
rels = d.part.rels
media = {rid: str(rel.target_ref).split("/")[-1] for rid, rel in rels.items() if "media" in str(rel.target_ref)}

paras = [p for p in body.findall(qn("w:p"))]
def ptext(p):
    return "".join(n.text or "" for n in p.iter(qn("w:t")))
def pimgs(p):
    return [media.get(b.get(qn("r:embed"))) for b in p.iter(qn("a:blip")) if b.get(qn("r:embed"))]

# 1) 每个图注段：紧邻的前一段必须含图片（图注在图片下方）
bad = []
for i, p in enumerate(paras):
    t = ptext(p).strip()
    if t.startswith(("图 3", "图 4", "图 5", "图 7")):
        prev = pimgs(paras[i - 1]) if i > 0 else []
        if not prev:
            bad.append((i, t[:30]))
print("captions without image above:", bad or "NONE")
# 2) 每个图片段：紧邻的后一段应是图注
bad2 = []
for i, p in enumerate(paras):
    if pimgs(p):
        nxt = ptext(paras[i + 1]).strip() if i + 1 < len(paras) else ""
        if not nxt.startswith(("图 3", "图 4", "图 5", "图 7")):
            bad2.append((i, pimgs(p), nxt[:30]))
print("images without caption below:", bad2 or "NONE")

# 3) 抽查：5.4 标题、参考文献尾部、目录条目、摘要
print("\n--- spot checks ---")
for i, p in enumerate(paras):
    t = ptext(p).strip()
    if t.startswith(("5.4 系统操作流程", "5.5 单模块展示", "5.6 数字仿真", "5.7 技术边界", "5.8 本章小结")):
        print(f"P{i}: {t[:40]}")
    if t.startswith("[9] Unity") or t.startswith("[19] Rundel") or t.startswith("[28] Rus"):
        print(f"P{i}: {t[:60]}")
    if t.startswith("5.4 系统操作流程18") or t.startswith("5.8 本章小结20"):
        print(f"P{i}: {t[:40]}")
    if t.startswith("报告另提出以正方形纸片"):
        print(f"P{i} 摘要: {t[:60]}...")

# 4) 字数
total = sum(len(ptext(p)) for p in paras)
tbl_chars = sum(len(n.text or "") for tb in d.tables for row in tb.rows for cell in row.cells for n in cell.paragraphs[0]._p.iter(qn("w:t")))
print(f"\n总字符数（正文段）: {total}；表格字符: {tbl_chars}；合计: {total + tbl_chars}")
